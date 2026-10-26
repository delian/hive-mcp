import io
import zipfile

import pytest
from docx import Document
from openpyxl import Workbook
from pptx import Presentation

from hive_mcp.domain.errors import ValidationError
from hive_mcp.ingest import office
from hive_mcp.ingest.extract import extract_text


def docx_bytes():
    doc = Document()
    doc.add_heading("Radar Plan", level=1)
    doc.add_paragraph("The pulse repetition frequency is 10 kHz.")
    doc.add_heading("Timing", level=2)
    p = doc.add_paragraph("First line")
    p.add_run().add_break()
    p.add_run("second line")
    table = doc.add_table(rows=2, cols=2)
    for r, row in enumerate((("Band", "Gain"), ("X", "35 dB"))):
        for c, value in enumerate(row):
            table.cell(r, c).text = value
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def xlsx_bytes():
    wb = Workbook()
    ws = wb.active
    ws.title = "Gains"
    ws.append(["Band", "Gain dB", "Note"])
    ws.append(["X", 35.5, "primary"])
    ws.append(["S", None, "spare"])
    ws2 = wb.create_sheet("Empty second")
    ws2["B2"] = "only B2"
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def pptx_bytes():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Range equation"
    slide.placeholders[1].text = "Power to the fourth root"
    slide.notes_slide.notes_text_frame.text = "mention the 19 percent figure"
    prs.slides.add_slide(prs.slide_layouts[5]).shapes.title.text = "Summary"
    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()


def test_docx_keeps_headings_paragraphs_breaks_and_tables():
    got = extract_text("plan.docx", docx_bytes())
    assert got.format == "docx"
    lines = got.text.splitlines()
    assert lines[0] == "# Radar Plan" and "## Timing" in lines
    assert "The pulse repetition frequency is 10 kHz." in lines
    assert "First line\nsecond line" in got.text
    assert "Band | Gain" in lines and "X | 35 dB" in lines


def test_xlsx_gives_every_sheet_with_names_and_aligned_columns():
    got = extract_text("gains.xlsx", xlsx_bytes())
    assert got.format == "xlsx"
    lines = got.text.splitlines()
    assert lines[0] == "## Gains" and "Band | Gain dB | Note" in lines
    assert "X | 35.5 | primary" in lines and "S |  | spare" in lines  # empty cell keeps its column
    assert "## Empty second" in lines and " | only B2" in lines


def test_pptx_gives_slides_titles_and_notes_in_order():
    got = extract_text("deck.pptx", pptx_bytes())
    lines = got.text.splitlines()
    assert lines.index("## Slide 1") < lines.index("## Slide 2")
    assert "### Range equation" in lines and "Power to the fourth root" in lines
    assert "Notes: mention the 19 percent figure" in lines and "### Summary" in lines


def test_office_title_extraction_survives_the_blank_file_check():
    for name, data in (("a.docx", Document),):
        out = io.BytesIO()
        data().save(out)
        with pytest.raises(ValidationError, match="no readable text"):
            extract_text(name, out.getvalue())


def zipped(files):
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for name, content in files.items():
            z.writestr(name, content)
    return out.getvalue()


def test_garbage_and_wrong_kinds_of_zip_are_refused_politely():
    for name, data in [
        ("x.docx", b"not a zip at all"),
        ("x.docx", zipped({"hello.txt": "hi"})),
        ("x.xlsx", zipped({"hello.txt": "hi"})),
        ("x.pptx", zipped({"hello.txt": "hi"})),
        ("x.docx", zipped({"word/document.xml": "<w:document><unclosed>"})),
    ]:
        with pytest.raises(ValidationError):
            extract_text(name, data)
    with pytest.raises(ValidationError, match="old binary"):
        extract_text("legacy.doc", b"\xd0\xcf\x11\xe0")


def test_xml_entity_bombs_are_refused_before_parsing():
    bomb = (
        '<?xml version="1.0"?><!DOCTYPE lolz [<!ENTITY a "aaaaaaaaaa"><!ENTITY b "&a;&a;&a;&a;">]>'
        '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        "<w:body><w:p><w:r><w:t>&b;</w:t></w:r></w:p></w:body></w:document>"
    )
    with pytest.raises(ValidationError, match="forbidden"):
        extract_text("bomb.docx", zipped({"word/document.xml": bomb}))
    lower = bomb.replace("DOCTYPE", "doctype")
    with pytest.raises(ValidationError, match="forbidden"):
        extract_text("bomb.docx", zipped({"word/document.xml": lower}))


def test_zip_bombs_are_refused_by_declared_sizes(monkeypatch):
    big = zipped({"word/document.xml": "x" * 1000})
    monkeypatch.setattr(office, "MAX_MEMBER_BYTES", 500)
    with pytest.raises(ValidationError, match="oversized"):
        extract_text("big.docx", big)
    monkeypatch.setattr(office, "MAX_MEMBER_BYTES", 10**9)
    monkeypatch.setattr(office, "MAX_TOTAL_BYTES", 500)
    with pytest.raises(ValidationError, match="too large"):
        extract_text("big.docx", big)
    monkeypatch.setattr(office, "MAX_TOTAL_BYTES", 10**9)
    monkeypatch.setattr(office, "MAX_MEMBERS", 1)
    with pytest.raises(ValidationError, match="too large"):
        extract_text("many.docx", zipped({"a": "1", "b": "2"}))


def test_spreadsheets_are_cut_at_the_row_limit(monkeypatch):
    monkeypatch.setattr(office, "MAX_ROWS", 3)
    wb = Workbook()
    for i in range(10):
        wb.active.append([f"r{i}"])
    out = io.BytesIO()
    wb.save(out)
    text = extract_text("long.xlsx", out.getvalue()).text
    assert "r2" in text and "r3" not in text and "more than 3 rows omitted" in text


def test_many_sheets_come_out_in_workbook_order_not_file_name_order():
    wb = Workbook()
    wb.active.title = "s1"
    wb.active["A1"] = "content 1"
    for i in range(2, 13):
        ws = wb.create_sheet(f"s{i}")
        ws["A1"] = f"content {i}"
    out = io.BytesIO()
    wb.save(out)
    headings = [
        ln
        for ln in extract_text("many.xlsx", out.getvalue()).text.splitlines()
        if ln.startswith("## ")
    ]
    assert headings == [f"## s{i}" for i in range(1, 13)]  # s10 does not sort before s2


def test_reordered_workbooks_follow_the_workbook_not_the_file_names():
    wb = Workbook()
    wb.active.title = "first"
    wb.active["A1"] = "one"
    wb.create_sheet("second")["A1"] = "two"
    wb.move_sheet("second", offset=-1)  # workbook order: second, first
    out = io.BytesIO()
    wb.save(out)
    text = extract_text("moved.xlsx", out.getvalue()).text
    assert text.index("## second") < text.index("## first")
    assert "## second\ntwo" in text and "## first\none" in text  # names stay attached to their data
