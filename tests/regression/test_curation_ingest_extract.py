"""Regression tests for text extraction (BUG-CUR-n)."""

from __future__ import annotations

import io
import time
import zipfile

import pytest

from hive_mcp.domain.errors import ValidationError
from hive_mcp.ingest.extract import extract_text

W = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
S = 'xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"'


def zipped(files: dict[str, str | bytes]) -> bytes:
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for name, content in files.items():
            z.writestr(name, content)
    return out.getvalue()


def test_html_without_a_closing_head_tag_keeps_the_body():
    """BUG-CUR-2: `</head>` is optional in HTML; the parser skipped everything after an unclosed
    <head>, so the page was stored as a lone '#'."""
    page = b"<html><head><title>T</title><body><h1>Hello</h1><p>World text</p></body></html>"
    got = extract_text("a.html", page)
    assert "World text" in got.text and got.title == "T"


def test_html_table_cells_do_not_run_together():
    """BUG-CUR-3: <td>Name</td><td>Age</td> was extracted as 'NameAge'."""
    page = b"<table><tr><td>Name</td><td>Age</td></tr><tr><td>Bob</td><td>7</td></tr></table>"
    assert extract_text("a.html", page).text == "Name | Age\nBob | 7"


def test_stray_nul_characters_do_not_make_a_text_file_unreadable():
    """BUG-CUR-4: a NUL after the first 8 KB (common in PDF text layers) reached the item
    constructor, which refused the whole document."""
    data = b"hello " * 3000 + b"\x00" + b"tail"
    assert extract_text("a.txt", data).text.endswith("hello tail")


def test_docx_reads_content_controls_and_nested_tables_once():
    """BUG-CUR-5: paragraphs inside <w:sdt> (content controls, tables of contents) were dropped
    and a nested table's text was repeated for every enclosing row."""
    doc = (
        f"<w:document {W}><w:body><w:p><w:r><w:t>Before</w:t></w:r></w:p>"
        "<w:sdt><w:sdtContent><w:p><w:r><w:t>Inside sdt</w:t></w:r></w:p></w:sdtContent></w:sdt>"
        "<w:tbl><w:tr><w:tc><w:p><w:r><w:t>outer</w:t></w:r></w:p><w:tbl><w:tr><w:tc><w:p><w:r>"
        "<w:t>inner</w:t></w:r></w:p></w:tc></w:tr></w:tbl></w:tc></w:tr></w:tbl>"
        "</w:body></w:document>"
    )
    text = extract_text("a.docx", zipped({"word/document.xml": doc})).text
    assert "Inside sdt" in text
    assert text.count("inner") == 1 and text.count("outer") == 1


def test_xlsx_cell_far_outside_the_sheet_is_refused_not_expanded():
    """BUG-CUR-6: a cell reference like ZZZZZZ1 made the reader emit a ~1 GB row (35 s of CPU)."""
    sheet = (
        f"<worksheet {S}><sheetData><row r='1'><c r='ZZZZZZ1' t='inlineStr'><is><t>x</t></is></c>"
        "</row></sheetData></worksheet>"
    )
    started = time.monotonic()
    with pytest.raises(ValidationError):
        extract_text("a.xlsx", zipped({"xl/worksheets/sheet1.xml": sheet}))
    assert time.monotonic() - started < 5


def test_xlsx_shared_string_amplification_is_capped():
    """BUG-CUR-7: one 1 MB shared string referenced from 20,000 rows produced 20 GB of text."""
    strings = f"<sst {S}><si><t>{'x' * 1_000_000}</t></si></sst>"
    rows = "".join(f"<row r='{i}'><c r='A{i}' t='s'><v>0</v></c></row>" for i in range(1, 3000))
    sheet = f"<worksheet {S}><sheetData>{rows}</sheetData></worksheet>"
    files = {"xl/worksheets/sheet1.xml": sheet, "xl/sharedStrings.xml": strings}
    with pytest.raises(ValidationError, match="too much text"):
        extract_text("a.xlsx", zipped(files))


def test_xlsx_cells_without_a_reference_keep_their_columns():
    """BUG-CUR-8: `r` is optional; every such cell was written to column A, losing all but one."""
    sheet = (
        f"<worksheet {S}><sheetData><row><c t='inlineStr'><is><t>a</t></is></c>"
        "<c t='inlineStr'><is><t>b</t></is></c></row></sheetData></worksheet>"
    )
    assert extract_text("a.xlsx", zipped({"xl/worksheets/sheet1.xml": sheet})).text.endswith(
        "a | b"
    )


def test_office_xml_with_a_doctype_is_refused_in_any_encoding():
    """BUG-CUR-9: the DTD check scanned bytes, so a UTF-16 document with an entity declaration
    was parsed (and its entities expanded)."""
    doc = (
        '<?xml version="1.0" encoding="UTF-16"?><!DOCTYPE d [<!ENTITY e "boom">]>'
        f"<w:document {W}><w:body><w:p><w:r><w:t>&e;</w:t></w:r></w:p></w:body></w:document>"
    ).encode("utf-16")
    with pytest.raises(ValidationError, match="forbidden"):
        extract_text("a.docx", zipped({"word/document.xml": doc}))


def test_pptx_table_text_is_extracted():
    """BUG-CUR-10: text inside slide tables (graphic frames) was dropped."""
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Sales"
    table = slide.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(4), Inches(1)).table
    for (r, c), text in {(0, 0): "Region", (0, 1): "Total", (1, 0): "EMEA", (1, 1): "42"}.items():
        table.cell(r, c).text = text
    out = io.BytesIO()
    prs.save(out)
    lines = extract_text("a.pptx", out.getvalue()).text.splitlines()
    assert "Region | Total" in lines and "EMEA | 42" in lines


def test_utf8_text_with_one_bad_byte_is_not_turned_into_mojibake():
    """BUG-CUR-17: any invalid byte made the whole file decode as Latin-1, so a Cyrillic or
    accented UTF-8 document became unreadable; legacy Latin-1 files must still work."""
    broken = ("Здравей свят, това е бележка. " * 20).encode() + b"\xff" + " край".encode()
    text = extract_text("a.txt", broken).text
    assert "Здравей свят" in text and "край" in text
    assert extract_text("l.txt", "café naïve".encode("latin-1")).text == "café naïve"


def test_lone_carriage_returns_are_line_breaks():
    """BUG-CUR-18: old-Mac line endings were kept as \\r, gluing lines and hiding headings."""
    assert extract_text("a.md", b"# Head\rbody\r").text == "# Head\nbody"


def test_a_heading_marker_alone_on_a_line_is_not_a_heading():
    """BUG-CUR-19: `\\s+` in the heading pattern crossed the line break, so '#' followed by the
    next line made that line a heading (and the title of the item)."""
    from hive_mcp.ingest.sections import split_sections

    (section,) = split_sections("#\nnext line\nbody")
    assert section.heading is None


async def test_titles_from_documents_are_one_clean_line(stack):
    """BUG-CUR-20: a multi-line <title> (or one with control characters) became the item title
    verbatim and one made only of zero-width characters made the whole ingest fail."""
    from tests.conformance.factory import ALICE

    page = b"<html><head><title>\n   Quarterly\n   plan\x07\n</title></head><body>text</body>"
    done = await stack.ingest.ingest(ALICE, "plan.html", page)
    assert done.items[0].title == "Quarterly plan"
    blank = "<title>\u200b</title><p>hello".encode()
    assert (await stack.ingest.ingest(ALICE, "blank.html", blank)).items[0].title == "blank.html"


def test_office_file_with_a_corrupt_member_is_refused_not_a_crash():
    """BUG-CUR-21: a zip member failing its CRC raised BadZipFile (not a HiveError) out of the
    extractor, which surfaced as an internal error instead of 'cannot read the file'."""
    doc = f"<w:document {W}><w:body><w:p><w:r><w:t>Hello there</w:t></w:r></w:p></w:body></w:document>"
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_STORED) as z:
        z.writestr("word/document.xml", doc)
    raw = bytearray(out.getvalue())
    at = raw.index(b"Hello there")
    raw[at : at + 5] = b"Jello"  # the stored CRC no longer matches
    with pytest.raises(ValidationError):
        extract_text("a.docx", bytes(raw))


def test_html_declaring_a_legacy_charset_is_decoded_with_it():
    """BUG-CUR-23: only UTF-8 and Latin-1 were tried, so a windows-1251 (or Shift_JIS) page
    turned into mojibake even though it says which encoding it uses."""
    page = '<html><head><meta charset="windows-1251"><title>Бележка</title></head>'
    page += "<body><p>Здравей, свят!</p></body></html>"
    got = extract_text("a.html", page.encode("cp1251"))
    assert got.title == "Бележка" and "Здравей, свят!" in got.text
    jp = (
        '<meta http-equiv="Content-Type" content="text/html; charset=Shift_JIS"><p>日本語のメモ</p>'
    )
    assert "日本語のメモ" in extract_text("j.html", jp.encode("shift_jis")).text


async def test_fetched_text_is_recoded_from_the_charset_of_the_response():
    """BUG-CUR-24: the Content-Type charset of a fetched page was ignored."""
    import httpx

    from hive_mcp.ingest.fetch import FetchPolicy, fetch

    body = "<p>Здравей, свят!</p>".encode("cp1251")

    def handler(request: httpx.Request) -> httpx.Response:
        return httpx.Response(
            200, content=body, headers={"content-type": "text/html; charset=windows-1251"}
        )

    async def resolve(host: str) -> list[str]:
        return ["93.184.216.34"]

    got = await fetch(
        "https://example.org/p",
        FetchPolicy(),
        resolve=resolve,
        transport=httpx.MockTransport(handler),
    )
    assert "Здравей, свят!" in extract_text(got.name, got.data).text


def test_xlsx_shared_strings_leave_out_the_phonetic_guide():
    """BUG-CUR-25: Excel stores furigana in <rPh> next to Japanese text; it was appended to the
    cell ('日本語ニホンゴ'), polluting search with text nobody typed."""
    strings = (
        f"<sst {S}><si><r><t>日本</t></r><r><t>語</t></r><rPh sb='0' eb='3'><t>ニホンゴ</t></rPh>"
        "<phoneticPr fontId='1'/></si></sst>"
    )
    sheet = (
        f"<worksheet {S}><sheetData><row r='1'><c r='A1' t='s'><v>0</v></c></row>"
        "</sheetData></worksheet>"
    )
    files = {"xl/worksheets/sheet1.xml": sheet, "xl/sharedStrings.xml": strings}
    text = extract_text("a.xlsx", zipped(files)).text
    assert "日本語" in text and "ニホンゴ" not in text
